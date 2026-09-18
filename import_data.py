"""
Script d'import : lit les fichiers .docx (tableaux Question | Réponse | Explication)
et le fichier .xlsx (Categorie | Question | BonneReponse), puis remplit questions.db.

Usage :
    python3 import_data.py fichier1.docx fichier2.docx ... fichier.xlsx
    python3 import_data.py --dossier /chemin/vers/dossier

Peut être relancé plusieurs fois : les questions déjà présentes (même texte
de question, comparé sans tenir compte des accents/majuscules) ne sont pas
dupliquées.
"""
import sqlite3
import sys
import os
import re
import unicodedata
import argparse

from assign_themes import choisir_theme

DB_PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)), "data", "questions.db")


def normalize(text: str) -> str:
    """Retire accents, ponctuation, espaces superflus, met en minuscules."""
    if text is None:
        return ""
    text = text.strip().lower()
    text = unicodedata.normalize("NFKD", text)
    text = "".join(c for c in text if not unicodedata.combining(c))
    text = re.sub(r"[^\w\s]", "", text)
    text = re.sub(r"\s+", " ", text).strip()
    return text


def init_db():
    os.makedirs(os.path.dirname(DB_PATH), exist_ok=True)
    conn = sqlite3.connect(DB_PATH)
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS questions (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            categorie TEXT,
            question TEXT NOT NULL,
            reponse TEXT NOT NULL,
            explication TEXT,
            question_norm TEXT NOT NULL,
            theme TEXT,
            fois_posee INTEGER DEFAULT 0,
            fois_reussie INTEGER DEFAULT 0
        )
        """
    )
    conn.execute(
        "CREATE UNIQUE INDEX IF NOT EXISTS idx_question_norm ON questions(question_norm)"
    )
    conn.commit()
    return conn


def category_from_filename(path: str) -> str:
    name = os.path.splitext(os.path.basename(path))[0]
    name = re.sub(r"^\d+_questions?_?", "", name, flags=re.IGNORECASE)
    name = name.replace("_", " ").strip(" _-")
    name = re.sub(r"\s+\d+\s*questions?$", "", name, flags=re.IGNORECASE)
    return name.strip().capitalize() or "Divers"


def import_docx(path: str, conn) -> int:
    import docx

    doc = docx.Document(path)
    category = category_from_filename(path)
    added = 0
    for table in doc.tables:
        rows = table.rows
        if not rows:
            continue
        header = [c.text.strip().lower() for c in rows[0].cells]
        start = 1 if any("question" in h for h in header) else 0
        for row in rows[start:]:
            cells = [c.text.strip() for c in row.cells]
            if len(cells) < 2:
                continue
            question = re.sub(r"^\d+[\.\)]\s*", "", cells[0]).strip()
            reponse = cells[1].strip()
            explication = cells[2].strip() if len(cells) > 2 else ""
            if not question or not reponse:
                continue
            added += insert_question(conn, category, question, reponse, explication)
    return added


def import_xlsx(path: str, conn) -> int:
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    added = 0
    for ws in wb.worksheets:
        rows = ws.iter_rows(values_only=True)
        header = next(rows, None)
        if not header:
            continue
        header_lower = [str(h).strip().lower() if h else "" for h in header]

        def col(*names):
            for n in names:
                if n in header_lower:
                    return header_lower.index(n)
            return None

        i_cat = col("categorie", "catégorie", "category")
        i_q = col("question")
        i_r = col("bonnereponse", "bonne reponse", "bonne réponse", "reponse", "réponse")

        if i_q is None or i_r is None:
            continue

        for row in rows:
            if row is None:
                continue
            question = str(row[i_q]).strip() if i_q < len(row) and row[i_q] else ""
            reponse = str(row[i_r]).strip() if i_r < len(row) and row[i_r] else ""
            categorie = (
                str(row[i_cat]).strip() if i_cat is not None and i_cat < len(row) and row[i_cat] else "Divers"
            )
            if not question or not reponse:
                continue
            added += insert_question(conn, categorie, question, reponse, "")
    return added


def insert_question(conn, categorie, question, reponse, explication) -> int:
    q_norm = normalize(question)
    if not q_norm:
        return 0
    theme = choisir_theme(categorie)
    try:
        conn.execute(
            "INSERT INTO questions (categorie, question, reponse, explication, question_norm, theme) "
            "VALUES (?, ?, ?, ?, ?, ?)",
            (categorie, question, reponse, explication, q_norm, theme),
        )
        return 1
    except sqlite3.IntegrityError:
        return 0  # question déjà présente


def import_file(path: str, conn) -> int:
    ext = os.path.splitext(path)[1].lower()
    if ext == ".docx":
        return import_docx(path, conn)
    elif ext in (".xlsx", ".xlsm"):
        return import_xlsx(path, conn)
    else:
        print(f"  (ignoré, extension non supportée : {path})")
        return 0


def main():
    parser = argparse.ArgumentParser(description="Importe des questions dans la base du bot.")
    parser.add_argument("fichiers", nargs="*", help="Fichiers .docx ou .xlsx à importer")
    parser.add_argument("--dossier", help="Importe tous les .docx/.xlsx d'un dossier")
    args = parser.parse_args()

    fichiers = list(args.fichiers)
    if args.dossier:
        for fn in os.listdir(args.dossier):
            if fn.lower().endswith((".docx", ".xlsx", ".xlsm")):
                fichiers.append(os.path.join(args.dossier, fn))

    if not fichiers:
        print("Aucun fichier fourni. Utilise --dossier ou passe les fichiers en argument.")
        sys.exit(1)

    conn = init_db()
    total = 0
    for f in fichiers:
        if not os.path.exists(f):
            print(f"  Fichier introuvable : {f}")
            continue
        added = import_file(f, conn)
        conn.commit()
        print(f"  {os.path.basename(f)} -> {added} nouvelles questions ajoutées")
        total += added

    count = conn.execute("SELECT COUNT(*) FROM questions").fetchone()[0]
    conn.close()
    print(f"\nTotal ajouté cette fois-ci : {total}")
    print(f"Total dans la base : {count} questions")


if __name__ == "__main__":
    main()
